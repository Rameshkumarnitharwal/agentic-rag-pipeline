"""PPTX loader using python-pptx — one Section per slide with text, tables, chart data and notes."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.shapes.base import BaseShape

from .base import Section, clean_text, rows_to_markdown

_TITLE_PLACEHOLDERS = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}


def _chart_block(shape: BaseShape) -> str:
    """Native charts store their data in the file, so extract it as a table."""
    chart = shape.chart
    title = chart.chart_title.text_frame.text.strip() if chart.has_title else ""
    plot = chart.plots[0] if len(chart.plots) else None
    if plot is None:
        return f"Chart: {title}" if title else ""
    categories = [str(c) for c in plot.categories]
    rows = [["Category", *[s.name for s in plot.series]]]
    for i, category in enumerate(categories):
        rows.append([category, *[_fmt(s.values[i]) if i < len(s.values) else "" for s in plot.series]])
    header = f"Chart ({chart.chart_type}): {title}".rstrip(": ")
    return f"{header}\n\n{rows_to_markdown(rows)}"


def _fmt(value: float | None) -> str:
    return "" if value is None else f"{value:g}"


def _shape_blocks(shape: BaseShape) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [b for child in shape.shapes for b in _shape_blocks(child)]
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
        return [rows_to_markdown(rows)]
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return [_chart_block(shape)]
    if shape.has_text_frame:
        lines = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
        return ["\n".join(lines)] if lines else []
    return []


def _slide_title(slide) -> str | None:
    for shape in slide.placeholders:
        if shape.placeholder_format.type in _TITLE_PLACEHOLDERS and shape.text_frame.text.strip():
            return shape.text_frame.text.strip()
    return None


def load_pptx(path: Path) -> list[Section]:
    presentation = Presentation(str(path))
    sections: list[Section] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        title = _slide_title(slide)
        blocks: list[str] = []
        for shape in slide.shapes:
            for block in _shape_blocks(shape):
                if block and block != title:
                    blocks.append(block)
        if title:
            blocks.insert(0, f"## {title}")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                blocks.append(f"Speaker notes: {notes}")

        text = clean_text("\n\n".join(blocks))
        if text:
            metadata = {"slide": slide_number}
            if title:
                metadata["slide_title"] = title
            sections.append(Section(text=text, metadata=metadata))
    return sections
